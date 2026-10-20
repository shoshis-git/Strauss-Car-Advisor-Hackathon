"""Rules, car catalogue, occupations, sample employees and reference calculations.

Everything is read from the data files (rules.json + challenge_data.xlsx) into
validated Pydantic models and exposed through a single `AppConfig` object.
Nothing business-related is hardcoded here - only the mapping of the Hebrew
sheet headers to field names (the file schema).
"""
from __future__ import annotations

import json
import re
from functools import lru_cache
from pathlib import Path
from typing import Any, Optional

import openpyxl
from pydantic import BaseModel, ConfigDict, Field, field_validator, model_validator
from pydantic.alias_generators import to_camel

from . import paths


# --------------------------------------------------------------------------- #
# rules.json
# --------------------------------------------------------------------------- #
class _RulesModel(BaseModel):
    """camelCase in JSON, snake_case in Python; unknown fields are preserved."""
    model_config = ConfigDict(alias_generator=to_camel, populate_by_name=True,
                              extra="allow", frozen=True)


class ElectricRule(_RulesModel):
    net_bonus: float
    method: str
    note: Optional[str] = None


class FamilyTrackRule(_RulesModel):
    min_children: int
    child_max_age: int
    cars: list[str]
    price_source: str
    gap_split_employee: float = Field(ge=0, le=1)
    note: Optional[str] = None


class WaiverRule(_RulesModel):
    allowed_by: str
    maintenance_is_taxable: bool
    maintenance_is_pensionable: bool


class TaxBracket(_RulesModel):
    up_to: Optional[float]  # None = no upper limit
    rate: float


class ContributionTable(_RulesModel):
    threshold: float
    ceiling: float
    low_rate: float
    high_rate: float


class PensionCreditRule(_RulesModel):
    rate: float
    max_contribution_rate: float
    income_ceiling: float
    note: Optional[str] = None


class ParticipationRule(_RulesModel):
    treatment: str
    reduces_value_of_use: bool
    note: Optional[str] = None


class RoundingRule(_RulesModel):
    during: str
    display: str
    note: Optional[str] = None


class NetMethod(_RulesModel):
    basis: str
    tax_year: int
    credit_point_value: float
    brackets: list[TaxBracket]
    national_insurance: ContributionTable
    health: ContributionTable
    pension_credit: PensionCreditRule
    contributions_base: str
    contributions_note: Optional[str] = None
    participation: ParticipationRule
    rounding: RoundingRule
    note: Optional[str] = None

    @field_validator("brackets", mode="before")
    @classmethod
    def _pairs_to_brackets(cls, v: Any) -> Any:
        return [{"up_to": b[0], "rate": b[1]} if isinstance(b, (list, tuple)) else b
                for b in v]

    @field_validator("brackets")
    @classmethod
    def _check_brackets(cls, v: list[TaxBracket]) -> list[TaxBracket]:
        limits = [b.up_to for b in v]
        if not v or limits[-1] is not None or None in limits[:-1]:
            raise ValueError("only the last tax bracket may (and must) be open-ended")
        if limits[:-1] != sorted(limits[:-1]):
            raise ValueError("tax brackets must be in ascending order")
        return v


class ValueOfUseRule(_RulesModel):
    source: str
    key: list[str]
    drive_field: dict[str, str]

    @property
    def drive_names(self) -> dict[int, str]:
        """Tax Authority drive code -> drive name (e.g. 3 -> 'חשמלי')."""
        return {int(k): v for k, v in self.drive_field.items()}

    def drive_code(self, drive_name: str) -> int:
        for code, name in self.drive_names.items():
            if name == drive_name:
                return code
        raise KeyError(f"unknown drive name: {drive_name!r}")


class Rules(_RulesModel):
    version: str
    note: Optional[str] = None
    budgets: dict[str, float]
    group_order: list[str]
    choose_from: str
    electric: ElectricRule
    family_track: FamilyTrackRule
    waiver: WaiverRule
    unused_budget: str
    net_method: NetMethod
    value_of_use: ValueOfUseRule

    @model_validator(mode="after")
    def _groups_consistent(self) -> "Rules":
        if set(self.group_order) != set(self.budgets):
            raise ValueError("groupOrder and budgets must list the same groups")
        return self

    def budget(self, group: str) -> float:
        return self.budgets[group]

    def group_rank(self, group: str) -> int:
        return self.group_order.index(group)


def load_rules(path: Path = paths.RULES_JSON) -> Rules:
    with open(path, encoding="utf-8") as f:
        return Rules.model_validate(json.load(f))


# --------------------------------------------------------------------------- #
# challenge_data.xlsx
# --------------------------------------------------------------------------- #
SHEET_CARS = "קטלוג רכבים"
SHEET_OCCUPATIONS = "עיסוקים"
SHEET_EMPLOYEES = "עובדים לדוגמה"
SHEET_CALCULATIONS = "חישובים"

CAR_COLUMNS = {
    "מזהה": "car_id",
    "קבוצה": "group",
    "תוצר": "manufacturer_name",
    "קוד תוצר": "manufacturer_code",
    "דגם (כינוי מסחרי)": "model_name",
    "קוד דגם": "model_code",
    "תאור דגם": "model_description",
    "שנת רישום": "registration_year",
    "הנעה": "drive",
    "מחיר מחירון": "list_price",
    "שווי שימוש חודשי": "value_of_use_monthly",
    "עלות חודשית לחברה (דמה)": "company_monthly_cost",
    "השתתפות עובד ברירת מחדל (דמה)": "default_employee_participation",
}
OCCUPATION_COLUMNS = {
    "עיסוק": "name",
    "קבוצת זכאות": "eligibility_group",
    "רכב חובה": "car_required",
    "מותר לוותר": "waiver_allowed",
    "אחזקת רכב בוויתור (₪)": "waiver_maintenance",
}
EMPLOYEE_COLUMNS = {
    "מזהה": "employee_id",
    "שם": "name",
    "ת״ז (פיקטיבית)": "national_id",
    "עיסוק": "occupation",
    "קבוצה": "group",
    "שכר בסיס": "base_salary",
    "עמלות": "commissions",
    "בונוס": "bonus",
    "פנסיה עובד %": "pension_employee_pct",
    "קרן השתלמות עובד %": "study_fund_employee_pct",
    "נקודות זיכוי": "credit_points",
    "ילדים עד 23": "children_under_23",
    "רכב היום": "current_car_raw",
    "השתתפות היום": "current_participation",
    "השאלה": "question",
}
CALCULATION_COLUMNS = {
    "מזהה": "employee_id",
    "שם": "name",
    "היום": "current_label",
    "החלופה": "alternative_label",
    "זכאות": "eligible",
    "סיבה": "reason",
    "נטו היום": "net_current",
    "נטו בחלופה": "net_alternative",
    "הפרש": "net_diff",
    "מס היום": "tax_current",
    "מס בחלופה": "tax_alternative",
    "ביטוח לאומי+בריאות היום": "ni_health_current",
    "ביטוח לאומי+בריאות בחלופה": "ni_health_alternative",
    "החזר חשמל ברוטו": "electric_refund_gross",
    "חלק העובד במסלול המשפחתי": "family_track_employee_share",
    "למה": "explanation",
}

_YES_NO = {"כן": True, "לא": False}
_MAINTENANCE_RE = re.compile(r"אחזקה\s*([\d,]+)")


def _yes_no(v: Any) -> Any:
    if isinstance(v, str):
        if v.strip() not in _YES_NO:
            raise ValueError(f"expected כן/לא, got {v!r}")
        return _YES_NO[v.strip()]
    return v


class _SheetModel(BaseModel):
    model_config = ConfigDict(frozen=True, str_strip_whitespace=True)


class Car(_SheetModel):
    car_id: str
    group: str
    manufacturer_name: str
    manufacturer_code: int  # '00885' in the sheet -> 885
    model_name: str
    model_code: int
    model_description: str
    registration_year: int
    drive: str
    list_price: float
    value_of_use_monthly: float
    company_monthly_cost: float
    default_employee_participation: float


class Occupation(_SheetModel):
    name: str
    eligibility_group: str
    car_required: bool
    waiver_allowed: bool
    waiver_maintenance: float

    @field_validator("car_required", "waiver_allowed", mode="before")
    @classmethod
    def _parse_yes_no(cls, v: Any) -> Any:
        return _yes_no(v)


class Employee(_SheetModel):
    employee_id: str
    name: str
    national_id: str
    occupation: str
    group: str
    base_salary: float
    commissions: float
    bonus: float
    pension_employee_pct: float
    study_fund_employee_pct: float
    credit_points: float
    children_under_23: int
    current_car_raw: str
    current_participation: float
    question: str
    # Resolved by AppConfig from current_car_raw:
    current_car_id: Optional[str] = None
    current_maintenance: Optional[float] = None

    @field_validator("national_id", mode="before")
    @classmethod
    def _id_as_str(cls, v: Any) -> Any:
        return str(v) if isinstance(v, int) else v


class ReferenceCalculation(_SheetModel):
    employee_id: str
    name: str
    current_label: str
    alternative_label: str
    eligible: bool
    reason: str
    net_current: float
    net_alternative: float
    net_diff: float
    tax_current: float
    tax_alternative: float
    ni_health_current: float
    ni_health_alternative: float
    electric_refund_gross: float
    family_track_employee_share: float
    explanation: str

    @field_validator("eligible", mode="before")
    @classmethod
    def _parse_yes_no(cls, v: Any) -> Any:
        return _yes_no(v)


def _read_sheet(wb: openpyxl.Workbook, sheet: str, columns: dict[str, str]) -> list[dict]:
    """Read a header-row sheet into dicts keyed by English field names."""
    if sheet not in wb.sheetnames:
        raise ValueError(f"sheet {sheet!r} not found in workbook")
    rows = wb[sheet].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else None for h in next(rows)]
    missing = set(columns) - set(header)
    if missing:
        raise ValueError(f"sheet {sheet!r} is missing columns: {sorted(missing)}")
    records = []
    for row in rows:
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in row):
            continue
        rec = {columns[h]: v for h, v in zip(header, row) if h in columns}
        records.append(rec)
    return records


# --------------------------------------------------------------------------- #
# Unified configuration
# --------------------------------------------------------------------------- #
class AppConfig(BaseModel):
    model_config = ConfigDict(frozen=True)

    rules: Rules
    cars: list[Car]
    occupations: list[Occupation]
    employees: list[Employee]
    reference_calculations: list[ReferenceCalculation]
    car_aliases: dict[str, list[str]] = Field(default_factory=dict)  # car_id -> names
    car_brand_words: list[str] = Field(default_factory=list)  # e.g. BYD, טויוטה

    @model_validator(mode="after")
    def _cross_validate(self) -> "AppConfig":
        groups = set(self.rules.group_order)
        car_ids = {c.car_id for c in self.cars}
        seen: dict[str, str] = {c.model_name.casefold(): c.car_id for c in self.cars}
        for car_id, aliases in self.car_aliases.items():
            if car_id not in car_ids:
                raise ValueError(f"car_aliases: {car_id} not in catalogue")
            for a in aliases:
                owner = seen.setdefault(a.strip().casefold(), car_id)
                if owner != car_id:
                    raise ValueError(f"car_aliases: {a!r} used for both {owner} and {car_id}")
        occ = {o.name: o for o in self.occupations}
        if len(car_ids) != len(self.cars):
            raise ValueError("duplicate car ids in catalogue")
        if len(occ) != len(self.occupations):
            raise ValueError("duplicate occupations")
        for car_id in self.rules.family_track.cars:
            if car_id not in car_ids:
                raise ValueError(f"familyTrack car {car_id} not in catalogue")
        for c in self.cars:
            if c.group not in groups and c.car_id not in self.rules.family_track.cars:
                raise ValueError(f"{c.car_id}: unknown group {c.group!r}")
            self.rules.value_of_use.drive_code(c.drive)  # drive must be a known name
        for o in self.occupations:
            if o.eligibility_group not in groups:
                raise ValueError(f"occupation {o.name}: unknown group {o.eligibility_group!r}")
        emp_ids = set()
        for e in self.employees:
            if e.occupation not in occ:
                raise ValueError(f"{e.employee_id}: unknown occupation {e.occupation!r}")
            if e.group != occ[e.occupation].eligibility_group:
                raise ValueError(f"{e.employee_id}: group differs from occupation group")
            emp_ids.add(e.employee_id)
        for r in self.reference_calculations:
            if r.employee_id not in emp_ids:
                raise ValueError(f"reference calculation for unknown employee {r.employee_id}")
        return self

    # --- lookups ---------------------------------------------------------- #
    def car(self, car_id: str) -> Car:
        for c in self.cars:
            if c.car_id == car_id:
                return c
        raise KeyError(car_id)

    def car_by_model_name(self, model_name: str) -> Optional[Car]:
        key = model_name.strip().casefold()
        for c in self.cars:
            if c.model_name.casefold() == key:
                return c
        return None

    def car_names(self, car: Car) -> list[str]:
        """Model name + configured aliases (Hebrew and alternative spellings)."""
        return [car.model_name, *self.car_aliases.get(car.car_id, [])]

    def cars_in_group(self, group: str) -> list[Car]:
        return [c for c in self.cars if c.group == group]

    def family_track_cars(self) -> list[Car]:
        return [self.car(cid) for cid in self.rules.family_track.cars]

    def occupation(self, name: str) -> Occupation:
        for o in self.occupations:
            if o.name == name:
                return o
        raise KeyError(name)

    def employee(self, employee_id: str) -> Employee:
        for e in self.employees:
            if e.employee_id == employee_id:
                return e
        raise KeyError(employee_id)

    def reference_calculation(self, employee_id: str) -> ReferenceCalculation:
        for r in self.reference_calculations:
            if r.employee_id == employee_id:
                return r
        raise KeyError(employee_id)


def _resolve_current_car(emp: dict, cars: list[Car]) -> dict:
    """Fill current_car_id (car in catalogue) or current_maintenance (waiver)."""
    raw = str(emp["current_car_raw"]).strip()
    for c in cars:
        if c.model_name.casefold() == raw.casefold():
            return {**emp, "current_car_id": c.car_id}
    m = _MAINTENANCE_RE.search(raw)
    if m:
        return {**emp, "current_maintenance": float(m.group(1).replace(",", ""))}
    raise ValueError(f"{emp['employee_id']}: cannot resolve current car {raw!r}")


def load_car_aliases(path: Path = paths.CAR_ALIASES_JSON) -> dict[str, list[str]]:
    """Optional file; missing file = no aliases."""
    return _load_car_names(path)[0]


def _load_car_names(path: Path) -> tuple[dict[str, list[str]], list[str]]:
    """(aliases per car id, brand words) from car_aliases.json; missing file = none."""
    if not path.exists():
        return {}, []
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    aliases = {k: [a.strip() for a in v] for k, v in data["aliases"].items()}
    return aliases, [b.strip() for b in data.get("brand_words", [])]


def load_config(rules_path: Path = paths.RULES_JSON,
                xlsx_path: Path = paths.CHALLENGE_XLSX,
                aliases_path: Path = paths.CAR_ALIASES_JSON) -> AppConfig:
    rules = load_rules(rules_path)
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    try:
        cars = [Car.model_validate(r) for r in _read_sheet(wb, SHEET_CARS, CAR_COLUMNS)]
        occupations = _read_sheet(wb, SHEET_OCCUPATIONS, OCCUPATION_COLUMNS)
        employees = [_resolve_current_car(e, cars)
                     for e in _read_sheet(wb, SHEET_EMPLOYEES, EMPLOYEE_COLUMNS)]
        calculations = _read_sheet(wb, SHEET_CALCULATIONS, CALCULATION_COLUMNS)
    finally:
        wb.close()
    aliases, brand_words = _load_car_names(aliases_path)
    return AppConfig(rules=rules, cars=cars, occupations=occupations,
                     employees=employees, reference_calculations=calculations,
                     car_aliases=aliases, car_brand_words=brand_words)


@lru_cache(maxsize=1)
def get_config() -> AppConfig:
    """Process-wide cached configuration from the default data directory."""
    return load_config()


def export_unified_config(config: AppConfig, path: Path = paths.UNIFIED_CONFIG_JSON) -> Path:
    """Write the whole validated configuration to one clean JSON file."""
    path.write_text(
        json.dumps(config.model_dump(mode="json"), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return path
